# -*- coding: utf-8 -*-
"""
QUẢN LÝ HỌC ONLINE Ô TÔ  —  server nội bộ (Python 3.9+, SQLite, openpyxl)
Chạy:  CHAY_TOOL.bat   (hoặc: python app/server.py)
"""
import os, sys, io, re, json, shutil, sqlite3, threading, time, uuid, hashlib, subprocess
import datetime as dt, unicodedata, webbrowser, socket, traceback, urllib.request
from functools import lru_cache
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse, parse_qs, unquote

APP_NAME = "QUẢN LÝ HỌC ONLINE Ô TÔ"
APP_ID = "hoc-online-oto"
VERSION = "1.3"
VERSION_DATE = "20/10/2026"
SCHEMA_VERSION = 4
GITHUB_REPO = "HP-wukoong/QUAN_LY_LT_ONLINE"
GITHUB_BRANCH = "main"
MINE = "MINE"   # người chạy mặc định
BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BASE)
sys.path.append(os.path.join(BASE, "vendor"))          # openpyxl kèm sẵn (dùng khi máy chưa cài)
PORT = int(os.environ.get("HOC_ONLINE_PORT", "8768"))
DATA_DIR = os.path.join(ROOT, "data")
BACKUP_DIR = os.path.join(ROOT, "backups")
STATIC_DIR = os.path.join(BASE, "static")
DB_PATH = os.path.join(DATA_DIR, "hoc_online.db")
os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(BACKUP_DIR, exist_ok=True)

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

DB_LOCK = threading.RLock()
ST_CHUA, ST_DANG, ST_XONG = "CHUA", "DANG", "XONG"
ST_LABEL = {ST_CHUA: "CHƯA HỌC", ST_DANG: "ĐANG HỌC", ST_XONG: "ĐÃ HỌC XONG"}
ST_RANK = {ST_CHUA: 0, ST_DANG: 1, ST_XONG: 2}

# key: (mặc định, kiểu, min, max)
SETTINGS_SPEC = {
    "max_gio_ngay": ("8", "float", 1, 24),
    "tinh_ngay_loc": ("auto", "enum", ("auto", "0", "1"), None),
    "dem_phut_mon": ("10", "int", 0, 120),
    "ngay_nghi_tuan": ("", "weekdays", None, None),
    "ngay_le": ("", "dates", None, None),
    "phi_gv": ("300000", "money", 0, 100000000),
    "don_gia_nth": ("0", "money", 0, 100000000),
    "ngay_canh_bao_loc": ("3", "int", 1, 60),
    "ngay_canh_bao_chua_hoc": ("5", "int", 1, 365),
    "ngay_dung_gio": ("2", "int", 1, 30),
    "backup_extra_dir": ("", "text", None, None),
    "backup_interval_m": ("1", "int", 1, 1440),
    "backup_giu": ("10", "int", 3, 200),
    "tu_kiem_tra_cap_nhat": ("1", "enum", ("0", "1"), None),
    # mẫu tin nhắn Zalo – {so} {xung_ho} {ten} {tien} {ten_day_du}
    "mau_phi": ("Em gửi danh sách {so} học viên học online {xung_ho} {ten} chưa đóng phí ({tien} đ):", "text", None, None),
    "mau_bao": ("Em gửi danh sách {so} học viên học online {xung_ho} {ten} đã chạy đủ giờ các môn:", "text", None, None),
    "mau_cuoi": ("", "text", None, None),
    "anh_tieu_de_phi": ("NHẮC ĐÓNG PHÍ HỌC ONLINE", "text", None, None),
    "anh_tieu_de_bao": ("HỌC VIÊN ĐÃ CHẠY ĐỦ GIỜ ONLINE", "text", None, None),
    "anh_chan": ("Chúc {xung_ho} nhiều sức khỏe – Lái xe an toàn!", "text", None, None),
    "ten_don_vi": ("TRUNG TÂM ĐÀO TẠO LÁI XE", "text", None, None),
}
DEFAULT_SETTINGS = {k: v[0] for k, v in SETTINGS_SPEC.items()}

# ----------------------------------------------------------------- tiện ích
def nfc(s):
    return unicodedata.normalize("NFC", str(s)) if s is not None else ""

def clean(s):
    return re.sub(r"\s+", " ", nfc(s)).strip()

def upper(s):
    return clean(s).upper()

def strip_accents(s):
    s = nfc(s).replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")

def norm_cccd(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\D", "", str(v))
    if len(s) == 11:              # mất số 0 đầu do Excel lưu dạng số
        s = "0" + s
    return s

def norm_khoa(s):
    s = upper(s)
    if not s:
        return ""
    s = re.sub(r"\.(XLSX|XLS|XLSM)$", "", s)
    s = re.sub(r"\s*\(\d+\)\s*$", "", s)                   # "BSTĐ-25 (1)" do trình duyệt tải trùng
    s = re.sub(r"\s*-\s*", "-", s).replace(" ", "")
    s = re.sub(r"^BST[_D]?(?=-|$|\d)", "BSTĐ", s)
    s = s.replace("BSTDD", "BSTĐ")
    m = re.match(r"^([A-ZĐ0-9.]+?)-?(\d+)$", s)
    if m and not re.match(r"^\d+$", m.group(1)):
        s = "%s-%02d" % (m.group(1).rstrip("-"), int(m.group(2)))
    return s

KHOA_RE = re.compile(r"^[A-ZĐ][A-ZĐ0-9.]{0,7}(-[A-ZĐ0-9.]{1,6})?-\d{1,3}$")

def person_key(s):
    """khóa so khớp tên người chạy: bỏ dấu, bỏ tiền tố xưng hô (A VUI = VUI)"""
    k = strip_accents(upper(s))
    k = re.sub(r"^(ANH|A|CHI|C|EM|E|CO|THAY|CHU|BAC)\.?\s+", "", k)
    return re.sub(r"[^A-Z0-9]", "", k)

def parse_money(v):
    s = re.sub(r"[^\d]", "", str(v or ""))
    return int(s) if s else 0

def parse_hm(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(round(float(v) * 60))
    if isinstance(v, dt.time):
        return v.hour * 60 + v.minute
    if isinstance(v, dt.timedelta):
        return int(v.total_seconds() // 60)
    s = str(v).strip()
    m = re.match(r"^(\d+)\s*[:hH]\s*(\d{1,2})", s)
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    try:
        return int(round(float(s.replace(",", ".")) * 60))
    except ValueError:
        return 0

def parse_pair(v):
    s = "" if v is None else str(v)
    if "/" in s:
        a, b = s.split("/", 1)
        return parse_hm(a), parse_hm(b)
    return parse_hm(s), 0

def fmt_hm(minutes):
    minutes = int(minutes or 0)
    return "%d:%02d" % (minutes // 60, minutes % 60)

def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()[:10]
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None

def today():
    return dt.date.today()

def now_s():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

def ddmm(s):
    d = to_date(s)
    return d.strftime("%d/%m/%Y") if d else ""

SUBJECT_SHORT = [("PHAP LUAT", "PL GTĐB"), ("KY THUAT LAI XE", "KTLX"), ("DAO DUC", "ĐĐ-VHGT"),
                 ("CAU TAO", "CT&SC"), ("KIEN THUC MOI", "KT MỚI"), ("MO PHONG", "MÔ PHỎNG")]

def subject_key(raw):
    s = re.sub(r"\(.*?\)", "", clean(raw))
    s = re.sub(r"\d{1,2}/\d{1,2}/\d{2,4}", "", s)
    return clean(s)

@lru_cache(maxsize=256)
def subject_short(name):
    a = strip_accents(name).upper()
    for k, v in SUBJECT_SHORT:
        if k in a:
            return v
    words = [w for w in re.split(r"[\s,&\-]+", name) if w]
    return "".join(w[0] for w in words[:5]).upper()

# ----------------------------------------------------------------- CSDL + nâng cấp
SCHEMA_V1 = """
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS teachers(id INTEGER PRIMARY KEY AUTOINCREMENT, ten TEXT UNIQUE NOT NULL, sdt TEXT DEFAULT '',
  ghi_chu TEXT DEFAULT '', active INTEGER DEFAULT 1, created_at TEXT);
CREATE TABLE IF NOT EXISTS runners(id INTEGER PRIMARY KEY AUTOINCREMENT, ten TEXT UNIQUE NOT NULL, sdt TEXT DEFAULT '',
  ghi_chu TEXT DEFAULT '', active INTEGER DEFAULT 1, created_at TEXT);
CREATE TABLE IF NOT EXISTS students(id INTEGER PRIMARY KEY AUTOINCREMENT,
  cccd TEXT UNIQUE NOT NULL, ho_ten TEXT NOT NULL, ngay_sinh TEXT DEFAULT '', ma_dk TEXT DEFAULT '', khoa TEXT DEFAULT '',
  hang TEXT DEFAULT '', gv_id INTEGER, nth_id INTEGER, phi_gv INTEGER DEFAULT 0, phi_gv_ngay TEXT DEFAULT '',
  nth_tt INTEGER DEFAULT 0, nth_tt_ngay TEXT DEFAULT '', trang_thai TEXT DEFAULT 'CHUA', trang_thai_ngay TEXT DEFAULT '',
  ngay_them TEXT, ngay_bat_dau TEXT DEFAULT '', ngay_xong TEXT DEFAULT '', ngay_loc TEXT DEFAULT '', import_id INTEGER,
  ghi_chu TEXT DEFAULT '', deleted INTEGER DEFAULT 0, updated_at TEXT);
CREATE INDEX IF NOT EXISTS ix_st_khoa ON students(khoa);
CREATE TABLE IF NOT EXISTS progress(student_id INTEGER, mon TEXT, mon_goc TEXT, thu_tu INTEGER,
  done_min INTEGER, norm_min INTEGER, dat INTEGER, ngay_loc TEXT, PRIMARY KEY(student_id, mon));
CREATE TABLE IF NOT EXISTS progress_history(id INTEGER PRIMARY KEY AUTOINCREMENT, student_id INTEGER, import_id INTEGER,
  ngay_loc TEXT, tong_done INTEGER, tong_norm INTEGER, con_lai INTEGER, chi_tiet TEXT, created_at TEXT);
CREATE INDEX IF NOT EXISTS ix_ph_st ON progress_history(student_id);
CREATE TABLE IF NOT EXISTS imports(id INTEGER PRIMARY KEY AUTOINCREMENT, ten_file TEXT, loai TEXT, khoa TEXT, ngay_loc TEXT,
  so_dong INTEGER, so_khop INTEGER, so_them INTEGER, so_xong INTEGER, ghi_chu TEXT, created_at TEXT);
CREATE TABLE IF NOT EXISTS logs(id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT, student_id INTEGER, hanh_dong TEXT, chi_tiet TEXT);
CREATE INDEX IF NOT EXISTS ix_log_st ON logs(student_id);
"""

MIGRATIONS = {
    2: [
        "ALTER TABLE students ADD COLUMN nguon TEXT DEFAULT ''",
        "ALTER TABLE students ADD COLUMN gio_loc TEXT DEFAULT ''",
        "ALTER TABLE students ADD COLUMN phi_gv_tien INTEGER DEFAULT 0",
        "ALTER TABLE students ADD COLUMN nth_tt_tien INTEGER DEFAULT 0",
        "ALTER TABLE teachers ADD COLUMN xung_ho TEXT DEFAULT 'Thầy'",
        "ALTER TABLE imports ADD COLUMN sha1 TEXT DEFAULT ''",
        "ALTER TABLE imports ADD COLUMN undone INTEGER DEFAULT 0",
        "CREATE TABLE IF NOT EXISTS import_undo(import_id INTEGER, student_id INTEGER, them INTEGER, snapshot TEXT)",
        "CREATE INDEX IF NOT EXISTS ix_undo ON import_undo(import_id)",
        "CREATE INDEX IF NOT EXISTS ix_ph_imp ON progress_history(import_id)",
        "DELETE FROM settings WHERE key IN ('backup_keep')",
        "UPDATE settings SET value='auto' WHERE key='tinh_ngay_loc' AND value NOT IN ('auto','0','1')",
    ],
    3: [
        "ALTER TABLE students ADD COLUMN bao_gv INTEGER DEFAULT 0",
        "ALTER TABLE students ADD COLUMN bao_gv_ngay TEXT DEFAULT ''",
        # HV đã xong từ file QL cũ (không có ngày xong) coi như đã báo GV từ trước
        "UPDATE students SET bao_gv=1 WHERE trang_thai='XONG' AND COALESCE(ngay_xong,'')=''",
    ],
    4: [
        "UPDATE teachers SET xung_ho='anh' WHERE COALESCE(xung_ho,'') IN ('','Thầy')",
        "UPDATE teachers SET xung_ho='chị' WHERE xung_ho='Cô'",
        "DELETE FROM settings WHERE key='backup_interval_h'",
    ],
}

def db():
    con = sqlite3.connect(DB_PATH, timeout=30)
    con.row_factory = sqlite3.Row
    return con

def migrate(con):
    v = con.execute("PRAGMA user_version").fetchone()[0]
    for ver in sorted(MIGRATIONS):
        if ver > v:
            for sql in MIGRATIONS[ver]:
                try:
                    con.execute(sql)
                except sqlite3.OperationalError as e:
                    if "duplicate column" not in str(e):
                        raise
            con.execute("PRAGMA user_version=%d" % ver)
    con.commit()

def init_db():
    with DB_LOCK:
        con = db()
        con.execute("PRAGMA journal_mode=WAL")
        con.executescript(SCHEMA_V1)
        need = con.execute("PRAGMA user_version").fetchone()[0] < SCHEMA_VERSION
        has_data = con.execute("SELECT COUNT(*) FROM students").fetchone()[0] > 0
        con.close()
        if need and has_data:
            do_backup("truoc_nang_cap")
        con = db()
        migrate(con)
        mine = mine_id(con)
        con.execute("UPDATE students SET nth_id=? WHERE nth_id IS NULL OR nth_id NOT IN (SELECT id FROM runners)", (mine,))
        for k, v in DEFAULT_SETTINGS.items():
            con.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", (k, v))
        con.commit()
        con.close()

def validate_setting(k, v):
    d, typ, lo, hi = SETTINGS_SPEC[k]
    s = str(v if v is not None else "").strip()
    if typ == "text":
        return s
    if typ == "enum":
        if s not in lo:
            raise ValueError("Giá trị không hợp lệ cho %s" % k)
        return s
    if typ == "money":
        n = parse_money(s)
        if n and n < 1000:
            raise ValueError("Số tiền %s đ quá nhỏ – nhập đủ số, vd 300000 hoặc 300.000" % s)
        return str(n)
    if typ in ("int", "float"):
        try:
            n = float(s.replace(",", "."))
        except ValueError:
            raise ValueError("'%s' không phải là số" % s)
        if not (lo <= n <= hi):
            raise ValueError("Giá trị %s phải từ %s đến %s" % (s, lo, hi))
        return str(int(round(n))) if typ == "int" else str(n)
    if typ == "weekdays":
        ds = "".join(sorted(set(re.findall(r"[0-6]", s))))
        if len(ds) >= 7:
            raise ValueError("Không thể nghỉ cả 7 ngày trong tuần")
        return ds
    if typ == "dates":
        out = []
        for x in re.split(r"[,;\s]+", s):
            if x:
                d0 = to_date(x)
                if not d0:
                    raise ValueError("Ngày lễ '%s' sai định dạng (YYYY-MM-DD hoặc dd/mm/yyyy)" % x)
                out.append(d0.isoformat())
        return ",".join(sorted(set(out)))
    return s

def get_settings(con=None):
    own = con is None
    con = con or db()
    s = dict(DEFAULT_SETTINGS)
    for r in con.execute("SELECT key,value FROM settings"):
        if r["key"] in SETTINGS_SPEC:
            try:
                s[r["key"]] = validate_setting(r["key"], r["value"])
            except ValueError:
                pass          # giá trị hỏng trong DB -> dùng mặc định, không làm sập tool
    if own:
        con.close()
    return s

def log(con, sid, act, detail=""):
    con.execute("INSERT INTO logs(ts,student_id,hanh_dong,chi_tiet) VALUES(?,?,?,?)", (now_s(), sid, act, detail))

XUNG_HO = ("anh", "chị", "thầy", "cô", "em")

def fill_tpl(tpl, **kw):
    out = tpl or ""
    for k, v in kw.items():
        out = out.replace("{%s}" % k, str(v))
    return out

def money_vn(n):
    return "{:,}".format(int(n or 0)).replace(",", ".")

def mine_id(con):
    """Người chạy mặc định MINE – HV chưa gán ai sẽ tự về MINE."""
    r = con.execute("SELECT id FROM runners WHERE ten=?", (MINE,)).fetchone()
    if r:
        return r["id"]
    return con.execute("INSERT INTO runners(ten,ghi_chu,created_at) VALUES(?,?,?)", (MINE, "Người chạy mặc định", now_s())).lastrowid

def find_or_create_teacher(con, name):
    name = upper(name)
    if not name:
        return None
    r = con.execute("SELECT id FROM teachers WHERE ten=?", (name,)).fetchone()
    if r:
        return r["id"]
    k = strip_accents(name)
    for r in con.execute("SELECT id,ten FROM teachers"):
        if strip_accents(r["ten"]) == k:
            return r["id"]
    return con.execute("INSERT INTO teachers(ten,created_at) VALUES(?,?)", (name, now_s())).lastrowid

def find_or_create_runner(con, name):
    name = upper(name)
    if not name:
        return None
    k = person_key(name)
    for r in con.execute("SELECT id,ten FROM runners"):
        if person_key(r["ten"]) == k:
            return r["id"]
    clean_name = re.sub(r"^(ANH|A|CHỊ|CHI|C|EM|E)\.?\s+", "", name) or name
    return con.execute("INSERT INTO runners(ten,created_at) VALUES(?,?)", (clean_name, now_s())).lastrowid

# ----------------------------------------------------------------- TIẾN ĐỘ & DỰ KIẾN
def mon_dat(m):
    """Đạt theo chữ trong file (kết luận của phần mềm học); ô trống mới xét giờ."""
    if m.get("dat") in (0, 1):
        return m["dat"]
    return 1 if m["norm_min"] and m["done_min"] >= m["norm_min"] else 0

def mon_lech(m):
    if not m["norm_min"]:
        return "định mức = 0"
    if m.get("dat") == 0 and m["done_min"] >= m["norm_min"]:
        return "đủ giờ nhưng file ghi Không Đạt"
    if m.get("dat") == 1 and m["done_min"] < m["norm_min"]:
        return "thiếu giờ nhưng file ghi Đạt"
    return ""

def calc_progress(rows, dem=0):
    tong_norm = sum(r["norm_min"] for r in rows)
    tong_done_cap = sum(min(r["done_min"], r["norm_min"]) for r in rows)
    tong_done = sum(r["done_min"] for r in rows)
    chua = [r for r in rows if not mon_dat(r)]
    con_lai = sum(max(0, r["norm_min"] - r["done_min"]) for r in chua)
    con_lai_dem = sum(max(0, r["norm_min"] - r["done_min"]) + dem for r in chua)
    lech = [(r["mon"], mon_lech(r)) for r in rows if mon_lech(r)]
    all_dat = bool(rows) and not chua and not any(not r["norm_min"] for r in rows)
    pct = round(tong_done_cap * 100.0 / tong_norm, 1) if tong_norm else 0.0
    if all_dat:
        pct, con_lai, con_lai_dem = 100.0, 0, 0
    elif pct >= 100:
        pct = 99.9
    return dict(tong_norm=tong_norm, tong_done=tong_done, tong_done_cap=tong_done_cap, con_lai=con_lai,
                con_lai_dem=con_lai_dem, all_dat=all_dat, any_done=any(r["done_min"] > 0 for r in rows), pct=pct,
                so_mon=len(rows), so_dat=len(rows) - len(chua), lech=lech)

class Ctx:
    """Tham số tính dự kiến – tạo 1 lần cho mỗi lượt tính."""
    def __init__(self, st):
        self.per_day = max(60, int(round(float(st["max_gio_ngay"]) * 60)))
        self.off = {int(x) for x in st["ngay_nghi_tuan"]}
        self.hol = {to_date(x) for x in st["ngay_le"].split(",") if x}
        self.mode = st["tinh_ngay_loc"]
        self.dem = int(st["dem_phut_mon"])
        self.today = today()

    def workday(self, d):
        return d.weekday() not in self.off and d not in self.hol

def forecast(con_lai, ngay_loc, gio_loc, c):
    """Trả về (số ngày chạy cần, ngày dự kiến xong). Lịch thật: bỏ ngày nghỉ/lễ.
    Ngày lọc – chế độ auto: nạp file trước 12h thì ngày lọc còn chạy được nửa ngày, sau 12h coi như hết."""
    if con_lai <= 0:
        return 0, None
    start = to_date(ngay_loc) or c.today
    if c.mode == "1":
        cap0 = c.per_day
    elif c.mode == "0":
        cap0 = 0
    else:
        cap0 = c.per_day // 2 if (gio_loc and gio_loc < "12:00") else 0
    rem, days, last = con_lai, 0, None
    if cap0 and c.workday(start):
        rem -= cap0
        days, last = 1, start
    d = start + dt.timedelta(days=1)
    guard = 0
    while rem > 0 and guard < 3000:
        if c.workday(d):
            rem -= c.per_day
            days += 1
            last = d
        d += dt.timedelta(days=1)
        guard += 1
    return days, last

# ----------------------------------------------------------------- ĐỌC FILE EXCEL
def load_wb(data):
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True)

def detect_type(wb):
    names = [strip_accents(n).upper() for n in wb.sheetnames]
    if any("HOC ONLINE" in n and "DANH SACH" in n for n in names):
        return "ql"
    ws = wb.worksheets[0]
    for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
        cells = [strip_accents(c).upper() for c in row if c]
        if any("TEN HOC VIEN" in c for c in cells) and any("MON HOC" in c for c in cells):
            return "report"
    return "unknown"

def parse_report(wb):
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = None
    for i, row in enumerate(rows[:6]):
        if any(c and "TEN HOC VIEN" in strip_accents(c).upper() for c in row):
            hdr_i = i
            break
    if hdr_i is None:
        raise ValueError("Không tìm thấy dòng tiêu đề 'Tên Học Viên'")
    hdr = [strip_accents(c).upper().strip() if c else "" for c in rows[hdr_i]]

    def col(*keys):
        for k in keys:
            for j, h in enumerate(hdr):
                if k in h:
                    return j
        return None
    c_ten, c_ns, c_hdt = col("TEN HOC VIEN"), col("NGAY SINH"), col("HANG DAO TAO")
    c_cccd, c_madk, c_mon = col("CAN CUOC", "CCCD", "CMND"), col("MA DANG KY"), col("MON HOC")
    if c_cccd is None or c_mon is None:
        raise ValueError("File thiếu cột Số Căn Cước hoặc Môn Học")
    sub_row = rows[hdr_i + 1] if len(rows) > hdr_i + 1 else ()
    subjects = []
    for j in range(c_mon, len(sub_row)):
        if sub_row[j]:
            raw = clean(sub_row[j])
            subjects.append(dict(col=j, raw=raw, key=subject_key(raw), short=subject_short(subject_key(raw))))
    if not subjects:
        raise ValueError("Không đọc được tên các môn học (dòng dưới tiêu đề 'Môn Học')")
    students, seen, trung = [], set(), []
    for row in rows[hdr_i + 2:]:
        if not row or c_ten >= len(row) or not row[c_ten]:
            continue
        cccd = norm_cccd(row[c_cccd] if c_cccd < len(row) else None)
        if not cccd:
            continue
        if cccd in seen:
            trung.append(cccd)
            continue
        seen.add(cccd)
        mons = []
        for k, sj in enumerate(subjects):
            v = row[sj["col"]] if sj["col"] < len(row) else None
            kq = row[sj["col"] + 1] if sj["col"] + 1 < len(row) else None
            done, norm = parse_pair(v)
            kqs = strip_accents(kq).upper().strip() if kq else ""
            dat = 1 if kqs == "DAT" else (0 if kqs else None)
            mons.append(dict(mon=sj["key"], mon_goc=sj["raw"], short=sj["short"], thu_tu=k,
                             done_min=done, norm_min=norm, dat=dat))
        nsd = to_date(row[c_ns]) if c_ns is not None and c_ns < len(row) else None
        students.append(dict(
            ho_ten=upper(row[c_ten]), cccd=cccd,
            ngay_sinh=nsd.strftime("%d/%m/%Y") if nsd else "",
            hang=clean(row[c_hdt]) if c_hdt is not None and c_hdt < len(row) else "",
            ma_dk=clean(row[c_madk]) if c_madk is not None and c_madk < len(row) else "",
            mons=mons))
    return dict(subjects=[dict(key=s["key"], short=s["short"], raw=s["raw"]) for s in subjects],
                students=students, trung=trung)

PAID_RE = re.compile(r"^(DA\s*(CK|TT|THANH\s*TOAN|CHUYEN|TRA)|OK|XONG)\b")

def parse_ql(wb):
    out = dict(teachers=[], students=[])
    for ws in wb.worksheets:
        n = strip_accents(ws.title).upper()
        if "GIAO VIEN" in n or n.startswith("DANH MUC GV"):
            for row in ws.iter_rows(min_row=1, values_only=True):
                if len(row) >= 2 and row[1] and isinstance(row[0], (int, float)):
                    out["teachers"].append(dict(ten=upper(row[1]), sdt=clean(row[2]) if len(row) > 2 and row[2] else ""))
    ws = next((w for w in wb.worksheets if "HOC ONLINE" in strip_accents(w.title).upper()
               and "DANH SACH" in strip_accents(w.title).upper()), None)
    if ws is None:
        return out
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next((i for i, row in enumerate(rows[:10]) if any(c and "HO VA TEN" in strip_accents(c).upper() for c in row)), None)
    if hdr_i is None:
        return out
    hdr = [strip_accents(c).upper().strip() if c else "" for c in rows[hdr_i]]

    def col(*keys):
        for k in keys:
            for j, h in enumerate(hdr):
                if k in h:
                    return j
        return None
    c_ten, c_cccd, c_khoa, c_gv = col("HO VA TEN"), col("CCCD", "CAN CUOC"), col("KHOA"), col("GIAO VIEN")
    c_phi, c_tt, c_nth = col("PHI"), col("TINH TRANG", "TRANG THAI"), col("NGUOI THUC HIEN")
    last = max(x for x in (c_nth, c_tt, c_phi, c_gv, c_cccd) if x is not None)
    for row in rows[hdr_i + 1:]:
        if not row or c_ten is None or not row[c_ten] or c_cccd is None:
            continue
        cccd = norm_cccd(row[c_cccd])
        if not cccd:
            continue
        tt = strip_accents(row[c_tt]).upper() if c_tt is not None and row[c_tt] else ""
        st = ST_XONG if ("DA HOC" in tt or "XONG" in tt) else (ST_DANG if "DANG" in tt else ST_CHUA)
        notes, nth_tt = [], 0
        for j in range(last + 1, len(row)):
            v = row[j]
            if v is None or str(v).strip() == "":
                continue
            s = strip_accents(v).upper().strip()
            if PAID_RE.match(s):
                nth_tt = 1
                if s in ("DA CK", "DA TT", "OK", "XONG", "DA THANH TOAN"):
                    continue
            notes.append(clean(v))
        phi_s = strip_accents(row[c_phi]).upper().strip() if c_phi is not None and row[c_phi] else ""
        out["students"].append(dict(
            ho_ten=upper(row[c_ten]), cccd=cccd, khoa=norm_khoa(row[c_khoa]) if c_khoa is not None else "",
            gv=upper(row[c_gv]) if c_gv is not None and row[c_gv] else "",
            phi_gv=1 if phi_s in ("XONG", "X", "OK", "DA DONG", "DA", "ROI", "DA CK", "DA TT") else 0,
            trang_thai=st, nth=upper(row[c_nth]) if c_nth is not None and row[c_nth] else "",
            nth_tt=nth_tt, ghi_chu=" | ".join(notes)))
    return out

# ----------------------------------------------------------------- XEM TRƯỚC / NẠP
PREVIEW = {}
PREVIEW_LOCK = threading.Lock()

def strip_upload_prefix(fn):
    return re.sub(r"^[0-9a-f]{8}-", "", os.path.basename(fn), flags=re.I)

def guess_khoa(filename, votes):
    k = norm_khoa(re.sub(r"\.(xlsx|xlsm|xls)$", "", strip_upload_prefix(filename), flags=re.I))
    if KHOA_RE.match(k or ""):
        return k
    return max(votes.items(), key=lambda x: x[1])[0] if votes else ""

def preview_file(filename, data):
    try:
        wb = load_wb(data)
    except Exception as e:
        raise ValueError("Không đọc được file Excel (%s). Nếu là .xls hãy mở bằng Excel và Lưu thành .xlsx" % e)
    kind = detect_type(wb)
    sha1 = hashlib.sha1(data).hexdigest()
    token = uuid.uuid4().hex
    con = db()
    try:
        dup = con.execute("SELECT created_at, ten_file FROM imports WHERE sha1=? AND undone=0 ORDER BY id DESC", (sha1,)).fetchone()
        dup_txt = ("File y hệt này đã nạp lúc %s (%s)" % (dup["created_at"], dup["ten_file"])) if dup else ""
        if kind == "report":
            rep = parse_report(wb)
            dem = int(get_settings(con)["dem_phut_mon"])
            known = {r["cccd"]: dict(r) for r in con.execute(
                "SELECT s.*, t.ten AS gv_ten FROM students s LEFT JOIN teachers t ON t.id=s.gv_id WHERE s.deleted=0")}
            by_madk = {r["ma_dk"]: r for r in known.values() if r.get("ma_dk")}
            cur_prog = {}
            for r in con.execute("SELECT student_id, mon, done_min FROM progress"):
                cur_prog.setdefault(r["student_id"], {})[r["mon"]] = r["done_min"]
            votes, items = {}, []
            for s in rep["students"]:
                k = known.get(s["cccd"]) or (by_madk.get(s["ma_dk"]) if s["ma_dk"] else None)
                p = calc_progress(s["mons"], dem)
                it = dict(s, pct=p["pct"], con_lai=p["con_lai"], all_dat=p["all_dat"], so_dat=p["so_dat"],
                          so_mon=p["so_mon"], lech=["%s: %s" % (subject_short(a), b) for a, b in p["lech"]])
                if k:
                    if k["khoa"]:
                        votes[k["khoa"]] = votes.get(k["khoa"], 0) + 1
                    old = k["trang_thai"]
                    new = ST_XONG if p["all_dat"] else (ST_DANG if p["any_done"] and old == ST_CHUA else old)
                    if ST_RANK[new] < ST_RANK[old]:
                        new = old
                    cp = cur_prog.get(k["id"], {})
                    giam = ["%s %s→%s" % (subject_short(m["mon"]), fmt_hm(cp[m["mon"]]), fmt_hm(m["done_min"]))
                            for m in s["mons"] if m["mon"] in cp and m["done_min"] < cp[m["mon"]]]
                    it.update(match=True, sid=k["id"], khoa_cu=k["khoa"], gv_ten=k.get("gv_ten") or "",
                              tt_cu=old, tt_moi=new, giam=giam,
                              ten_khac=k["ho_ten"] if strip_accents(k["ho_ten"]) != strip_accents(s["ho_ten"]) else "",
                              canh_bao=("Đang ghi ĐÃ HỌC XONG nhưng file chưa đạt đủ" if old == ST_XONG and not p["all_dat"] else ""))
                else:
                    it.update(match=False)
                items.append(it)
            khoa = guess_khoa(filename, votes)
            thieu = []
            if khoa:
                in_file = {s["cccd"] for s in rep["students"]}
                thieu = [dict(r) for r in con.execute(
                    "SELECT id,ho_ten,cccd,trang_thai FROM students WHERE deleted=0 AND khoa=? AND trang_thai<>'XONG'", (khoa,))
                    if r["cccd"] not in in_file]
            with PREVIEW_LOCK:
                PREVIEW[token] = dict(kind=kind, filename=filename, rep=rep, sha1=sha1, created=time.time())
            return dict(token=token, kind=kind, filename=strip_upload_prefix(filename), khoa=khoa,
                        subjects=rep["subjects"], items=items, so_dong=len(items), so_khop=sum(1 for i in items if i["match"]),
                        thieu=thieu, trung=rep["trung"], dup=dup_txt)
        if kind == "ql":
            q = parse_ql(wb)
            known = {r["cccd"] for r in con.execute("SELECT cccd FROM students")}
            runners = {person_key(r["ten"]) for r in con.execute("SELECT ten FROM runners")}
            gop = {}
            for s in q["students"]:
                if s["nth"] and person_key(s["nth"]) not in runners:
                    gop.setdefault(person_key(s["nth"]), set()).add(s["nth"])
            with PREVIEW_LOCK:
                PREVIEW[token] = dict(kind=kind, filename=filename, ql=q, sha1=sha1, created=time.time())
            return dict(token=token, kind=kind, filename=strip_upload_prefix(filename), so_gv=len(q["teachers"]),
                        so_hv=len(q["students"]), so_trung=sum(1 for s in q["students"] if s["cccd"] in known), dup=dup_txt,
                        nguoi_chay_moi=[" = ".join(sorted(v)) for v in gop.values()])
        raise ValueError("Không nhận dạng được loại file. Hỗ trợ: Learning Report (xuất từ phần mềm học online) "
                         "hoặc file quản lý QL_HOC_ONLINE.xlsx.")
    finally:
        con.close()

def take_preview(tok, kinds):
    with PREVIEW_LOCK:
        pv = PREVIEW.pop(tok, None)          # lấy ra ngay -> bấm đúp không nạp 2 lần
    if not pv or pv["kind"] not in kinds:
        raise ValueError("File này đã được nạp hoặc phiên xem trước hết hạn – hãy chọn lại file.")
    return pv

SNAP_FIELDS = ("trang_thai", "trang_thai_ngay", "ngay_loc", "gio_loc", "ngay_bat_dau", "ngay_xong", "import_id",
               "khoa", "ma_dk", "hang", "ngay_sinh", "deleted")

def snapshot(con, sid):
    s = con.execute("SELECT %s FROM students WHERE id=?" % ",".join(SNAP_FIELDS), (sid,)).fetchone()
    prog = [dict(r) for r in con.execute("SELECT * FROM progress WHERE student_id=?", (sid,))]
    return json.dumps(dict(st=dict(s), prog=prog), ensure_ascii=False)

def write_history(con, sid, imp, ngay_loc, mons, p):
    chi = json.dumps([[m["thu_tu"], m["done_min"], m["norm_min"], mon_dat(m)] for m in mons])
    last = con.execute("SELECT id, ngay_loc FROM progress_history WHERE student_id=? ORDER BY ngay_loc DESC, id DESC LIMIT 1", (sid,)).fetchone()
    if last and last["ngay_loc"] == ngay_loc:      # lọc lại trong cùng ngày -> thay dòng của ngày đó
        con.execute("UPDATE progress_history SET import_id=?, tong_done=?, tong_norm=?, con_lai=?, chi_tiet=?, created_at=? WHERE id=?",
                    (imp, p["tong_done_cap"], p["tong_norm"], p["con_lai"], chi, now_s(), last["id"]))
    else:
        con.execute("INSERT INTO progress_history(student_id,import_id,ngay_loc,tong_done,tong_norm,con_lai,chi_tiet,created_at) VALUES(?,?,?,?,?,?,?,?)",
                    (sid, imp, ngay_loc, p["tong_done_cap"], p["tong_norm"], p["con_lai"], chi, now_s()))

def apply_progress(con, sid, mons, ngay_loc, gio_loc, imp, dem):
    cur = con.execute("SELECT trang_thai, ngay_loc, ngay_bat_dau, ngay_xong FROM students WHERE id=?", (sid,)).fetchone()
    p = calc_progress(mons, dem)
    if cur["ngay_loc"] and ngay_loc < cur["ngay_loc"]:
        return cur["trang_thai"], cur["trang_thai"], "cu"           # dữ liệu cũ hơn: không ghi đè
    con.execute("DELETE FROM progress WHERE student_id=?", (sid,))
    for m in mons:
        con.execute("INSERT INTO progress(student_id,mon,mon_goc,thu_tu,done_min,norm_min,dat,ngay_loc) VALUES(?,?,?,?,?,?,?,?)",
                    (sid, m["mon"], m["mon_goc"], m["thu_tu"], m["done_min"], m["norm_min"], m["dat"], ngay_loc))
    write_history(con, sid, imp, ngay_loc, mons, p)
    old = cur["trang_thai"]
    new = ST_XONG if p["all_dat"] else (ST_DANG if p["any_done"] and old == ST_CHUA else old)
    if ST_RANK[new] < ST_RANK[old]:
        new = old                                                   # không bao giờ tự lùi trạng thái
    sets = {"ngay_loc": ngay_loc, "gio_loc": gio_loc, "import_id": imp, "updated_at": now_s()}
    if p["any_done"] and not cur["ngay_bat_dau"]:
        sets["ngay_bat_dau"] = ngay_loc
    if new != old:
        sets.update(trang_thai=new, trang_thai_ngay=now_s())
        log(con, sid, "TỰ ĐỔI TRẠNG THÁI", "%s → %s (lọc ngày %s)" % (ST_LABEL[old], ST_LABEL[new], ddmm(ngay_loc)))
    if new == ST_XONG and old != ST_XONG and not cur["ngay_xong"]:
        sets["ngay_xong"] = ngay_loc
    con.execute("UPDATE students SET %s WHERE id=?" % ",".join("%s=?" % k for k in sets), (*sets.values(), sid))
    return old, new, "ok"

def commit_report(body):
    ngay = to_date(body.get("ngay_loc")) or today()
    if ngay > today():
        raise ValueError("Ngày lọc %s ở tương lai – kiểm tra lại" % ngay.strftime("%d/%m/%Y"))
    pv = take_preview(body.get("token"), ("report",))
    ngay_loc = ngay.isoformat()
    gio_loc = dt.datetime.now().strftime("%H:%M") if ngay == today() else "23:59"
    khoa = norm_khoa(body.get("khoa") or "")
    add = {norm_cccd(a["cccd"]): a for a in body.get("them", [])}
    doi_khoa = bool(body.get("doi_khoa"))
    cho_giam = bool(body.get("cho_giam"))
    dem = int(get_settings()["dem_phut_mon"])
    with DB_LOCK:
        con = db()
        try:
            cur_prog = {}
            for r in con.execute("SELECT student_id, mon, done_min FROM progress"):
                cur_prog.setdefault(r["student_id"], {})[r["mon"]] = r["done_min"]
            known = {r["cccd"]: r["id"] for r in con.execute("SELECT id,cccd FROM students WHERE deleted=0")}
            deleted = {r["cccd"]: r["id"] for r in con.execute("SELECT id,cccd FROM students WHERE deleted=1")}
            by_madk = {r["ma_dk"]: r["id"] for r in con.execute("SELECT id,ma_dk FROM students WHERE deleted=0 AND ma_dk<>''")}
            imp = con.execute("INSERT INTO imports(ten_file,loai,khoa,ngay_loc,so_dong,created_at,sha1) VALUES(?,?,?,?,?,?,?)",
                              (strip_upload_prefix(pv["filename"]), "report", khoa, ngay_loc, len(pv["rep"]["students"]), now_s(), pv["sha1"])).lastrowid
            n_khop = n_them = n_xong = n_cu = n_giam = 0
            for s in pv["rep"]["students"]:
                sid = known.get(s["cccd"]) or (by_madk.get(s["ma_dk"]) if s["ma_dk"] else None)
                them = 0
                if not sid and s["cccd"] in add:
                    a = add[s["cccd"]]
                    gv_id = int(a["gv_id"]) if a.get("gv_id") else None
                    nth_id = int(a["nth_id"]) if a.get("nth_id") else mine_id(con)
                    if s["cccd"] in deleted:
                        sid = deleted[s["cccd"]]
                        con.execute("INSERT INTO import_undo VALUES(?,?,?,?)", (imp, sid, 2, snapshot(con, sid)))
                        con.execute("UPDATE students SET deleted=0, ho_ten=?, khoa=?, gv_id=?, nth_id=?, trang_thai='CHUA', "
                                    "ngay_them=?, updated_at=? WHERE id=?", (s["ho_ten"], khoa, gv_id, nth_id, ngay_loc, now_s(), sid))
                    else:
                        sid = con.execute(
                            "INSERT INTO students(cccd,ho_ten,ngay_sinh,ma_dk,khoa,hang,gv_id,nth_id,trang_thai,ngay_them,nguon,updated_at)"
                            " VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", (s["cccd"], s["ho_ten"], s["ngay_sinh"], s["ma_dk"], khoa, s["hang"],
                                                                 gv_id, nth_id, ST_CHUA, ngay_loc, "file tiến độ", now_s())).lastrowid
                        con.execute("INSERT INTO import_undo VALUES(?,?,?,?)", (imp, sid, 1, ""))
                    log(con, sid, "THÊM TỪ FILE", strip_upload_prefix(pv["filename"]))
                    n_them += 1
                    them = 1
                elif sid:
                    cp = cur_prog.get(sid, {})
                    if not cho_giam and any(m["mon"] in cp and m["done_min"] < cp[m["mon"]] for m in s["mons"]):
                        n_giam += 1
                        continue
                    con.execute("INSERT INTO import_undo VALUES(?,?,?,?)", (imp, sid, 0, snapshot(con, sid)))
                    n_khop += 1
                else:
                    continue
                r = con.execute("SELECT ngay_sinh,ma_dk,hang,khoa FROM students WHERE id=?", (sid,)).fetchone()
                upd = {}
                if not r["ngay_sinh"] and s["ngay_sinh"]: upd["ngay_sinh"] = s["ngay_sinh"]
                if s["ma_dk"] and r["ma_dk"] != s["ma_dk"]: upd["ma_dk"] = s["ma_dk"]
                if s["hang"] and r["hang"] != s["hang"]: upd["hang"] = s["hang"]
                if khoa and not them and (not r["khoa"] or (doi_khoa and r["khoa"] != khoa)):
                    upd["khoa"] = khoa
                    if r["khoa"]:
                        log(con, sid, "ĐỔI KHÓA", "%s → %s" % (r["khoa"], khoa))
                if upd:
                    con.execute("UPDATE students SET %s WHERE id=?" % ",".join("%s=?" % k for k in upd), (*upd.values(), sid))
                old, new, flag = apply_progress(con, sid, s["mons"], ngay_loc, gio_loc, imp, dem)
                n_cu += flag == "cu"
                n_xong += (new == ST_XONG and old != ST_XONG)
            note = []
            if n_cu: note.append("%d HV đã có dữ liệu ngày mới hơn – bỏ qua" % n_cu)
            if n_giam: note.append("%d HV bị giảm giờ – bỏ qua" % n_giam)
            con.execute("UPDATE imports SET so_khop=?, so_them=?, so_xong=?, ghi_chu=? WHERE id=?",
                        (n_khop, n_them, n_xong, "; ".join(note), imp))
            con.commit()
        finally:
            con.close()
    return dict(ok=True, so_khop=n_khop, so_them=n_them, so_xong=n_xong, so_cu=n_cu, so_giam=n_giam, import_id=imp)

def undo_import(imp_id):
    with DB_LOCK:
        con = db()
        try:
            last = con.execute("SELECT id FROM imports WHERE loai='report' AND undone=0 ORDER BY id DESC LIMIT 1").fetchone()
            if not last or last["id"] != imp_id:
                raise ValueError("Chỉ hoàn tác được lần nạp file tiến độ GẦN NHẤT (hoàn tác lần lượt từ mới về cũ).")
            n = 0
            for u in con.execute("SELECT * FROM import_undo WHERE import_id=?", (imp_id,)).fetchall():
                sid = u["student_id"]
                if u["them"] == 1:
                    con.execute("DELETE FROM progress WHERE student_id=?", (sid,))
                    con.execute("DELETE FROM students WHERE id=?", (sid,))
                else:
                    snap = json.loads(u["snapshot"])
                    st = snap["st"]
                    con.execute("UPDATE students SET %s WHERE id=?" % ",".join("%s=?" % k for k in st), (*st.values(), sid))
                    con.execute("DELETE FROM progress WHERE student_id=?", (sid,))
                    for p in snap["prog"]:
                        con.execute("INSERT INTO progress(%s) VALUES(%s)" % (",".join(p), ",".join("?" * len(p))), tuple(p.values()))
                    log(con, sid, "HOÀN TÁC NẠP FILE", "lần nạp #%d" % imp_id)
                n += 1
            con.execute("DELETE FROM progress_history WHERE import_id=?", (imp_id,))
            con.execute("UPDATE imports SET undone=1, ghi_chu=COALESCE(ghi_chu,'')||' [ĐÃ HOÀN TÁC '||?||']' WHERE id=?", (now_s(), imp_id))
            con.commit()
            return dict(ok=True, so=n)
        finally:
            con.close()

def commit_ql(body):
    pv = take_preview(body.get("token"), ("ql",))
    q = pv["ql"]
    phi_now = int(get_settings()["phi_gv"])
    n_gv = n_them = n_capnhat = n_bo = 0
    with DB_LOCK:
        con = db()
        try:
            for t in q["teachers"]:
                r = con.execute("SELECT id,sdt FROM teachers WHERE ten=?", (t["ten"],)).fetchone()
                if r:
                    if t["sdt"] and not r["sdt"]:
                        con.execute("UPDATE teachers SET sdt=? WHERE id=?", (t["sdt"], r["id"]))
                else:
                    con.execute("INSERT INTO teachers(ten,sdt,created_at) VALUES(?,?,?)", (t["ten"], t["sdt"], now_s()))
                    n_gv += 1
            imp = con.execute("INSERT INTO imports(ten_file,loai,so_dong,created_at,sha1) VALUES(?,?,?,?,?)",
                              (strip_upload_prefix(pv["filename"]), "ql", len(q["students"]), now_s(), pv["sha1"])).lastrowid
            for s in q["students"]:
                r = con.execute("SELECT * FROM students WHERE cccd=?", (s["cccd"],)).fetchone()
                if r:
                    if r["deleted"]:
                        n_bo += 1
                        continue
                    # CHỈ BỔ SUNG: điền ô trống, không lùi trạng thái, không bỏ tích, ghi chú nối thêm
                    upd, ch = {}, []
                    if not r["gv_id"] and s["gv"]:
                        upd["gv_id"] = find_or_create_teacher(con, s["gv"]); ch.append("GV: (trống) → %s" % s["gv"])
                    if (not r["nth_id"] or r["nth_id"] == mine_id(con)) and s["nth"]:
                        upd["nth_id"] = find_or_create_runner(con, s["nth"]); ch.append("Người chạy: (trống) → %s" % s["nth"])
                    if not r["khoa"] and s["khoa"]:
                        upd["khoa"] = s["khoa"]; ch.append("Khóa: (trống) → %s" % s["khoa"])
                    if s["phi_gv"] and not r["phi_gv"]:
                        upd["phi_gv"] = 1; upd["phi_gv_tien"] = phi_now; upd["phi_gv_ngay"] = today().isoformat(); ch.append("Phí GV: chưa → đã đóng")
                    if s["nth_tt"] and not r["nth_tt"]:
                        upd["nth_tt"] = 1; ch.append("TT người chạy: chưa → đã")
                    if ST_RANK[s["trang_thai"]] > ST_RANK[r["trang_thai"]]:
                        upd["trang_thai"] = s["trang_thai"]; ch.append("Trạng thái: %s → %s" % (ST_LABEL[r["trang_thai"]], ST_LABEL[s["trang_thai"]]))
                    if s["ghi_chu"] and s["ghi_chu"] not in (r["ghi_chu"] or ""):
                        upd["ghi_chu"] = ((r["ghi_chu"] + " | ") if r["ghi_chu"] else "") + s["ghi_chu"]; ch.append("Ghi chú: + " + s["ghi_chu"])
                    if upd:
                        upd["updated_at"] = now_s()
                        con.execute("UPDATE students SET %s WHERE id=?" % ",".join("%s=?" % k for k in upd), (*upd.values(), r["id"]))
                        log(con, r["id"], "BỔ SUNG TỪ FILE QL", "; ".join(ch))
                        n_capnhat += 1
                    else:
                        n_bo += 1
                    continue
                vals = dict(cccd=s["cccd"], ho_ten=s["ho_ten"], khoa=s["khoa"],
                            gv_id=find_or_create_teacher(con, s["gv"]) if s["gv"] else None,
                            nth_id=find_or_create_runner(con, s["nth"]) if s["nth"] else mine_id(con),
                            phi_gv=s["phi_gv"], phi_gv_tien=phi_now if s["phi_gv"] else 0, nth_tt=s["nth_tt"], trang_thai=s["trang_thai"], ghi_chu=s["ghi_chu"],
                            bao_gv=1 if s["trang_thai"] == ST_XONG else 0,
                            ngay_them="", nguon="file QL cũ", deleted=0, updated_at=now_s())
                sid = con.execute("INSERT INTO students(%s) VALUES(%s)" % (",".join(vals), ",".join("?" * len(vals))),
                                  tuple(vals.values())).lastrowid
                log(con, sid, "NHẬP TỪ FILE QL", strip_upload_prefix(pv["filename"]))
                n_them += 1
            con.execute("UPDATE imports SET so_them=?, so_khop=?, ghi_chu=? WHERE id=?",
                        (n_them, n_capnhat, "GV mới: %d, không đổi: %d" % (n_gv, n_bo), imp))
            con.commit()
        finally:
            con.close()
    return dict(ok=True, so_gv=n_gv, so_them=n_them, so_capnhat=n_capnhat, so_bo=n_bo)

# ----------------------------------------------------------------- DANH SÁCH (1 lần tính, dùng chung)
def compute_all(include_deleted=False):
    con = db()
    try:
        st = get_settings(con)
        c = Ctx(st)
        prog = {}
        for r in con.execute("SELECT student_id, mon, done_min, norm_min, dat, thu_tu FROM progress ORDER BY student_id, thu_tu"):
            prog.setdefault(r["student_id"], []).append(dict(r))
        hist = {}
        for r in con.execute("SELECT student_id, ngay_loc, tong_done FROM progress_history ORDER BY student_id, ngay_loc"):
            hist.setdefault(r["student_id"], []).append((r["ngay_loc"], r["tong_done"]))
        rows = con.execute("""SELECT s.*, t.ten gv_ten, t.sdt gv_sdt, t.xung_ho gv_xh, n.ten nth_ten FROM students s
            LEFT JOIN teachers t ON t.id=s.gv_id LEFT JOIN runners n ON n.id=s.nth_id
            WHERE s.deleted=? ORDER BY s.id""", (1 if include_deleted else 0,)).fetchall()
    finally:
        con.close()
    since = (c.today - dt.timedelta(days=14)).isoformat()
    nloc = int(st["ngay_canh_bao_loc"])
    out = []
    for r in rows:
        d = dict(r)
        mons = prog.get(r["id"], [])
        p = calc_progress(mons, c.dem)
        d["mons"] = [dict(short=subject_short(m["mon"]), mon=m["mon"], done=m["done_min"], norm=m["norm_min"], dat=mon_dat(m)) for m in mons]
        d.update(tong_done=p["tong_done_cap"], tong_norm=p["tong_norm"], con_lai=p["con_lai"], pct=p["pct"],
                 all_dat=p["all_dat"], so_mon=p["so_mon"], so_dat=p["so_dat"],
                 lech="; ".join("%s: %s" % (subject_short(a), b) for a, b in p["lech"]))
        d["mon_thieu"] = ", ".join("%s còn %s" % (subject_short(m["mon"]), fmt_hm(max(0, m["norm_min"] - m["done_min"])))
                                   for m in mons if not mon_dat(m))
        d["so_ngay"], d["ngay_du_kien"], d["con_ngay"], d["qua_du_kien"] = None, "", None, False
        if mons and r["trang_thai"] != ST_XONG and p["con_lai_dem"] > 0:
            n, ngay = forecast(p["con_lai_dem"], r["ngay_loc"], r["gio_loc"], c)
            d["so_ngay"] = n
            if ngay:
                d["ngay_du_kien"] = ngay.isoformat()
                d["con_ngay"] = max(0, (ngay - c.today).days)      # không bao giờ âm
                d["qua_du_kien"] = ngay < c.today
        d["so_ngay_tu_loc"] = (c.today - to_date(r["ngay_loc"])).days if r["ngay_loc"] else None
        # tốc độ thực tế (tham khảo) & số ngày đứng giờ
        h = hist.get(r["id"], [])
        d["toc_do"], d["dung_gio"] = None, 0
        if len(h) >= 2:
            recent = [x for x in h if x[0] >= since]
            a, b = (recent[0], recent[-1]) if len(recent) >= 2 else (h[-2], h[-1])
            days = (to_date(b[0]) - to_date(a[0])).days
            if days > 0:
                d["toc_do"] = round((b[1] - a[1]) / 60.0 / days, 1)
            j = len(h) - 1
            while j > 0 and h[j][1] <= h[j - 1][1]:
                j -= 1
            if j < len(h) - 1:
                d["dung_gio"] = (to_date(h[-1][0]) - to_date(h[j][0])).days
        d["can_loc"] = bool(r["trang_thai"] == ST_DANG and (not r["ngay_loc"] or d["so_ngay_tu_loc"] >= nloc or d["qua_du_kien"]))
        out.append(d)
    return out, st

def reminders(S, st):
    nchua = int(st["ngay_canh_bao_chua_hoc"])
    ndg = int(st["ngay_dung_gio"])
    td = today()
    R = []

    def add(muc, loai, s, msg):
        R.append(dict(muc=muc, loai=loai, id=s["id"], ho_ten=s["ho_ten"], khoa=s["khoa"], gv=s.get("gv_ten") or "",
                      nth=s.get("nth_ten") or "", msg=msg))
    for s in S:
        if s["trang_thai"] == ST_XONG:
            if not s["bao_gv"]:
                add(1, "CHƯA BÁO GV", s, "Đã chạy xong%s – chưa báo lại cho GV %s" % (
                    (" ngày " + ddmm(s["ngay_xong"])) if s["ngay_xong"] else "", s.get("gv_ten") or "(chưa có GV)"))
            if s["mons"] and not s["all_dat"]:
                add(1, "XONG CHƯA ĐẠT", s, "Đang ghi đã xong nhưng file lọc %s chưa đạt: %s" % (ddmm(s["ngay_loc"]), s["mon_thieu"] or s["lech"]))
            continue
        if s["lech"]:
            add(1, "LỆCH KẾT QUẢ", s, "Kiểm tra trên phần mềm học: " + s["lech"])
        if s["trang_thai"] == ST_DANG:
            if s["qua_du_kien"]:
                add(1, "LẼ RA ĐÃ XONG", s, "Theo %sh/ngày lẽ ra xong %s – nạp file mới để xác nhận" % (st["max_gio_ngay"].rstrip("0").rstrip("."), ddmm(s["ngay_du_kien"])))
            elif s["dung_gio"] >= ndg:
                add(1, "ĐỨNG GIỜ", s, "%d ngày không tăng giờ – kiểm tra người chạy" % s["dung_gio"])
            elif s["can_loc"]:
                add(2, "CẦN LỌC LẠI", s, ("Đã %d ngày chưa nạp file tiến độ" % s["so_ngay_tu_loc"]) if s["ngay_loc"] else "Chưa có dữ liệu tiến độ")
            elif s["ngay_du_kien"] and to_date(s["ngay_du_kien"]) <= td + dt.timedelta(days=1):
                add(3, "SẮP XONG", s, "Dự kiến xong %s – lọc để chốt" % ddmm(s["ngay_du_kien"]))
        if s["trang_thai"] == ST_CHUA:
            d0 = to_date(s["ngay_them"])
            if d0 and (td - d0).days >= nchua:
                add(2, "CHƯA HỌC LÂU", s, "Nhập %d ngày vẫn chưa chạy" % (td - d0).days)
    R.sort(key=lambda x: (x["muc"], x["loai"], x["khoa"], x["ho_ten"]))
    phi = int(st["phi_gv"])
    dg = int(st["don_gia_nth"])
    can_loc = {}
    for s in S:
        if s["can_loc"] or (s["trang_thai"] == ST_DANG and s["ngay_du_kien"] and to_date(s["ngay_du_kien"]) <= td + dt.timedelta(days=1)):
            x = can_loc.setdefault(s["khoa"] or "(chưa có khóa)", dict(khoa=s["khoa"] or "(chưa có khóa)", so=0, cu_nhat=None))
            x["so"] += 1
            if s["ngay_loc"] and (x["cu_nhat"] is None or s["ngay_loc"] < x["cu_nhat"]):
                x["cu_nhat"] = s["ngay_loc"]
    gom = {}
    for s in S:
        if not s["phi_gv"] and s.get("gv_ten"):
            gom.setdefault((s["gv_ten"], s.get("gv_sdt") or "", (s.get("gv_xh") or "anh").lower()), []).append(s)
    cuoi = st["mau_cuoi"].strip()

    def sec(title, g):
        return "%s (%d):\n%s" % (title, len(g), "\n".join("- %s (%s)" % (x["ho_ten"].title(), x["khoa"]) for x in g))
    tin = []
    for (gv, sdt, xh), lst in sorted(gom.items(), key=lambda x: -len(x[1])):
        ten_goi = gv.split()[-1].title()
        parts = []
        for st_k, title in ((ST_XONG, "Đã học xong"), (ST_DANG, "Đang chạy"), (ST_CHUA, "Chưa chạy")):
            g = [x for x in lst if x["trang_thai"] == st_k]
            if g:
                parts.append(sec(title, g))
        tien = len(lst) * phi
        kw = dict(so=len(lst), xung_ho=xh, ten=ten_goi, tien=money_vn(tien), ten_day_du=gv.title())
        text = fill_tpl(st["mau_phi"], **kw) + "\n\n" + "\n\n".join(parts) + (("\n\n" + fill_tpl(cuoi, **kw)) if cuoi else "")
        tin.append(dict(gv=gv, sdt=sdt, so=len(lst), tien=tien, so_xong=sum(1 for x in lst if x["trang_thai"] == ST_XONG),
                        loai="phi", xung_ho=xh, ten_goi=ten_goi, text=text))
    # báo GV: HV đã xong chưa báo, gom theo GV
    bao = {}
    for s in S:
        if s["trang_thai"] == ST_XONG and not s["bao_gv"]:
            k = (s.get("gv_ten") or "(chưa có GV)", s.get("gv_sdt") or "", (s.get("gv_xh") or "anh").lower())
            bao.setdefault(k, []).append(s)
    bao_gv = []
    for (gv, sdt, xh), lst in sorted(bao.items(), key=lambda x: -len(x[1])):
        ten_goi = gv.split()[-1].title()
        no_phi = [x for x in lst if not x["phi_gv"]]
        co_phi = [x for x in lst if x["phi_gv"]]
        parts = []
        if co_phi:
            parts.append(sec("Đã đóng phí", co_phi))
        if no_phi:
            parts.append(sec("Chưa đóng phí", no_phi))
        kw = dict(so=len(lst), xung_ho=xh, ten=ten_goi, tien=money_vn(len(no_phi) * phi), ten_day_du=gv.title())
        text = fill_tpl(st["mau_bao"], **kw) + "\n\n" + "\n\n".join(parts) + (("\n\n" + fill_tpl(cuoi, **kw)) if cuoi else "")
        bao_gv.append(dict(gv=gv, sdt=sdt, so=len(lst), ids=[x["id"] for x in lst], so_no_phi=len(no_phi),
                           loai="bao", xung_ho=xh, ten_goi=ten_goi, text=text))
    tra = {}
    for s in S:
        if s["trang_thai"] == ST_XONG and s["nth_id"] and not s["nth_tt"] and s.get("nth_ten") != MINE:
            x = tra.setdefault(s["nth_ten"], dict(ten=s["nth_ten"], so=0, tien=0, ids=[]))
            x["so"] += 1
            x["tien"] += dg
            x["ids"].append(s["id"])
    return dict(items=R, can_loc=sorted(can_loc.values(), key=lambda x: -x["so"]), tin_gv=tin, bao_gv=bao_gv,
                tra_nth=sorted(tra.values(), key=lambda x: -x["so"]))

def stats(S, st):
    phi = int(st["phi_gv"])
    dg = int(st["don_gia_nth"])

    def grp(key):
        g = {}
        for s in S:
            k = s.get(key) or "(chưa gán)"
            x = g.setdefault(k, dict(ten=k, tong=0, chua=0, dang=0, xong=0, phi_ok=0, phi_no=0, tien_thu=0, tt_ok=0, tt_no=0,
                                     tien_tra=0, toc=[], sdt=s.get("gv_sdt") if key == "gv_ten" else ""))
            x["tong"] += 1
            x[{"CHUA": "chua", "DANG": "dang", "XONG": "xong"}[s["trang_thai"]]] += 1
            if s["phi_gv"]:
                x["phi_ok"] += 1
                x["tien_thu"] += s["phi_gv_tien"] or phi
            else:
                x["phi_no"] += 1
            if s["nth_tt"]:
                x["tt_ok"] += 1
                x["tien_tra"] += s["nth_tt_tien"] or dg
            elif s["trang_thai"] == ST_XONG and s["nth_id"] and s.get("nth_ten") != MINE:
                x["tt_no"] += 1
            if s["trang_thai"] == ST_DANG and s["toc_do"] is not None:
                x["toc"].append(s["toc_do"])
        for x in g.values():
            x["tien_no"] = x["phi_no"] * phi
            x["tien_tra_no"] = x["tt_no"] * dg
            x["toc_do"] = round(sum(x["toc"]) / len(x["toc"]), 1) if x["toc"] else None
            del x["toc"]
        return sorted(g.values(), key=lambda x: (-x["tong"], x["ten"]))
    T = dict(tong=len(S), chua=sum(s["trang_thai"] == ST_CHUA for s in S), dang=sum(s["trang_thai"] == ST_DANG for s in S),
             xong=sum(s["trang_thai"] == ST_XONG for s in S), phi_ok=sum(s["phi_gv"] for s in S),
             tien_thu=sum((s["phi_gv_tien"] or phi) for s in S if s["phi_gv"]),
             tt_no=sum(1 for s in S if s["trang_thai"] == ST_XONG and s["nth_id"] and not s["nth_tt"] and s.get("nth_ten") != MINE),
             chua_bao=sum(1 for s in S if s["trang_thai"] == ST_XONG and not s["bao_gv"]),
             phi_gv=phi, don_gia_nth=dg)
    T["phi_no"] = T["tong"] - T["phi_ok"]
    T["tien_no"] = T["phi_no"] * phi
    T["tien_tra_no"] = T["tt_no"] * dg
    thang = {}

    def mrow(m):
        return thang.setdefault(m, dict(thang=m, xong=0, phi_so=0, phi_tien=0, tt_so=0, tt_tien=0, moi=0))
    for s in S:
        if s["ngay_xong"]: mrow(s["ngay_xong"][:7])["xong"] += 1
        if s["phi_gv"] and s["phi_gv_ngay"]:
            x = mrow(s["phi_gv_ngay"][:7]); x["phi_so"] += 1; x["phi_tien"] += s["phi_gv_tien"] or phi
        if s["nth_tt"] and s["nth_tt_ngay"]:
            x = mrow(s["nth_tt_ngay"][:7]); x["tt_so"] += 1; x["tt_tien"] += s["nth_tt_tien"] or dg
        if s["ngay_them"]: mrow(s["ngay_them"][:7])["moi"] += 1
    return dict(tong=T, gv=grp("gv_ten"), nth=grp("nth_ten"), khoa=grp("khoa"),
                thang=sorted(thang.values(), key=lambda x: x["thang"], reverse=True))

def student_detail(sid):
    con = db()
    try:
        s = con.execute("""SELECT s.*, t.ten gv_ten, t.sdt gv_sdt, n.ten nth_ten FROM students s
            LEFT JOIN teachers t ON t.id=s.gv_id LEFT JOIN runners n ON n.id=s.nth_id WHERE s.id=?""", (sid,)).fetchone()
        if not s:
            raise ValueError("Không tìm thấy học viên")
        hist = [dict(r) for r in con.execute("SELECT ngay_loc,tong_done,tong_norm,con_lai FROM progress_history WHERE student_id=? ORDER BY ngay_loc, id", (sid,))]
        logs = [dict(r) for r in con.execute("SELECT ts,hanh_dong,chi_tiet FROM logs WHERE student_id=? ORDER BY id DESC LIMIT 100", (sid,))]
    finally:
        con.close()
    return dict(student=dict(s), hist=hist, logs=logs)

# ----------------------------------------------------------------- SỬA / THÊM / XÓA
EDITABLE = {"ho_ten", "cccd", "ngay_sinh", "ma_dk", "khoa", "hang", "gv_id", "nth_id", "phi_gv", "nth_tt", "bao_gv", "trang_thai", "ghi_chu"}

def update_students(ids, fields, note=""):
    fields = {k: v for k, v in fields.items() if k in EDITABLE}
    if not fields or not ids:
        raise ValueError("Không có gì để cập nhật")
    if "trang_thai" in fields and fields["trang_thai"] not in ST_LABEL:
        raise ValueError("Trạng thái không hợp lệ")
    if "khoa" in fields: fields["khoa"] = norm_khoa(fields["khoa"])
    if "ho_ten" in fields:
        fields["ho_ten"] = upper(fields["ho_ten"])
        if not fields["ho_ten"]: raise ValueError("Họ tên không được để trống")
    if "cccd" in fields:
        fields["cccd"] = norm_cccd(fields["cccd"])
        if len(ids) > 1: raise ValueError("Không sửa CCCD hàng loạt")
        if len(fields["cccd"]) not in (9, 12): raise ValueError("CCCD phải có 12 số (hoặc CMND 9 số)")
    for k in ("ghi_chu", "ngay_sinh", "ma_dk", "hang"):
        if k in fields: fields[k] = clean(fields[k])[:500]
    for k in ("gv_id", "nth_id"):
        if k in fields: fields[k] = int(fields[k]) if fields[k] not in (None, "", 0, "0") else None
    for k in ("phi_gv", "nth_tt", "bao_gv"):
        if k in fields: fields[k] = 1 if fields[k] in (1, True, "1", "true") else 0
    st = get_settings()
    with DB_LOCK:
        con = db()
        try:
            changed = 0
            for sid in ids:
                old = con.execute("SELECT * FROM students WHERE id=?", (sid,)).fetchone()
                if not old:
                    continue
                f = dict(fields)
                if "phi_gv" in f and f["phi_gv"] != old["phi_gv"]:
                    f["phi_gv_ngay"] = today().isoformat() if f["phi_gv"] else ""
                    f["phi_gv_tien"] = int(st["phi_gv"]) if f["phi_gv"] else 0
                if "bao_gv" in f and f["bao_gv"] != old["bao_gv"]:
                    f["bao_gv_ngay"] = today().isoformat() if f["bao_gv"] else ""
                if "nth_id" in f and not f["nth_id"]:
                    f["nth_id"] = mine_id(con)
                if "nth_tt" in f and f["nth_tt"] != old["nth_tt"]:
                    f["nth_tt_ngay"] = today().isoformat() if f["nth_tt"] else ""
                    f["nth_tt_tien"] = int(st["don_gia_nth"]) if f["nth_tt"] else 0
                if "trang_thai" in f and f["trang_thai"] != old["trang_thai"]:
                    f["trang_thai_ngay"] = now_s()
                    if f["trang_thai"] == ST_DANG and not old["ngay_bat_dau"]:
                        f["ngay_bat_dau"] = today().isoformat()
                    if f["trang_thai"] == ST_XONG and not old["ngay_xong"]:
                        f["ngay_xong"] = today().isoformat()
                    if f["trang_thai"] != ST_XONG:
                        f["ngay_xong"] = ""
                changes = ["%s: %s → %s" % (k, old[k], v) for k, v in fields.items()
                           if (old[k] if old[k] is not None else "") != (v if v is not None else "")]
                if not changes:
                    continue
                f["updated_at"] = now_s()
                try:
                    con.execute("UPDATE students SET %s WHERE id=?" % ",".join("%s=?" % k for k in f), (*f.values(), sid))
                except sqlite3.IntegrityError:
                    raise ValueError("CCCD đã có ở học viên khác")
                log(con, sid, "SỬA" + (" (" + note + ")" if note else ""), "; ".join(changes))
                changed += 1
            con.commit()
        finally:
            con.close()
    return dict(ok=True, so=changed)

def add_students(lst, common):
    if not lst:
        raise ValueError("Danh sách trống")
    them, bo = 0, []
    st = get_settings()
    with DB_LOCK:
        con = db()
        try:
            for b in lst:
                cccd, ten = norm_cccd(b.get("cccd")), upper(b.get("ho_ten"))
                if not ten or len(cccd) not in (9, 12):
                    bo.append("%s %s: thiếu tên hoặc CCCD không đủ 12 số" % (ten, cccd)); continue
                r = con.execute("SELECT id,deleted,ho_ten FROM students WHERE cccd=?", (cccd,)).fetchone()
                if r and not r["deleted"]:
                    bo.append("%s: đã có trong danh sách (%s)" % (cccd, r["ho_ten"])); continue
                g = dict(common, **{k: v for k, v in b.items() if v not in (None, "")})
                vals = dict(ho_ten=ten, khoa=norm_khoa(g.get("khoa")), hang=clean(g.get("hang")),
                            gv_id=int(g["gv_id"]) if g.get("gv_id") else None, nth_id=int(g["nth_id"]) if g.get("nth_id") else mine_id(con),
                            phi_gv=1 if g.get("phi_gv") else 0, ghi_chu=clean(g.get("ghi_chu")), ma_dk=clean(g.get("ma_dk")),
                            ngay_sinh=clean(g.get("ngay_sinh")), trang_thai=ST_CHUA, ngay_them=today().isoformat(),
                            nguon="thêm tay", deleted=0, updated_at=now_s(), ngay_loc="", ngay_xong="", ngay_bat_dau="")
                if vals["phi_gv"]:
                    vals["phi_gv_ngay"] = today().isoformat()
                    vals["phi_gv_tien"] = int(st["phi_gv"])
                if r:
                    con.execute("DELETE FROM progress WHERE student_id=?", (r["id"],))
                    con.execute("UPDATE students SET %s WHERE id=?" % ",".join("%s=?" % k for k in vals), (*vals.values(), r["id"]))
                    sid = r["id"]
                else:
                    vals["cccd"] = cccd
                    sid = con.execute("INSERT INTO students(%s) VALUES(%s)" % (",".join(vals), ",".join("?" * len(vals))), tuple(vals.values())).lastrowid
                log(con, sid, "THÊM TAY", "")
                them += 1
            con.commit()
        finally:
            con.close()
    return dict(ok=True, them=them, bo=bo)

def delete_students(ids, restore=False):
    with DB_LOCK:
        con = db()
        for sid in ids:
            con.execute("UPDATE students SET deleted=?, updated_at=? WHERE id=?", (0 if restore else 1, now_s(), sid))
            log(con, sid, "KHÔI PHỤC" if restore else "XÓA (vào thùng rác)", "")
        con.commit()
        con.close()
    return dict(ok=True)

def purge_students(ids):
    with DB_LOCK:
        con = db()
        n = 0
        for sid in ids:
            r = con.execute("SELECT cccd, ho_ten FROM students WHERE id=? AND deleted=1", (sid,)).fetchone()
            if not r:
                continue                                     # chỉ xóa vĩnh viễn HV đang trong thùng rác
            for t in ("progress", "progress_history", "import_undo"):
                con.execute("DELETE FROM %s WHERE student_id=?" % t, (sid,))
            con.execute("DELETE FROM students WHERE id=?", (sid,))
            log(con, None, "XÓA VĨNH VIỄN", "%s - %s (id %d)" % (r["ho_ten"], r["cccd"], sid))
            n += 1
        con.commit()
        con.close()
    return dict(ok=True, so=n)

# ----------------------------------------------------------------- DANH MỤC
def list_cat(table):
    con = db()
    col = "gv_id" if table == "teachers" else "nth_id"
    rows = [dict(r) for r in con.execute(
        "SELECT c.*, (SELECT COUNT(*) FROM students s WHERE s.%s=c.id AND s.deleted=0) so_hv FROM %s c ORDER BY c.active DESC, c.ten" % (col, table))]
    con.close()
    keyf = person_key if table == "runners" else strip_accents
    keys = {}
    for r in rows:
        keys.setdefault(keyf(r["ten"]), []).append(r["id"])
    for r in rows:
        others = [i for i in keys[keyf(r["ten"])] if i != r["id"]]
        r["trung_voi"] = others[0] if others else None
    return rows

def save_cat(table, b):
    ten = upper(b.get("ten"))
    if not ten:
        raise ValueError("Tên không được trống")
    col = "gv_id" if table == "teachers" else "nth_id"
    with DB_LOCK:
        con = db()
        try:
            cid = int(b["id"]) if b.get("id") else None
            if cid and b.get("gop_vao"):
                tgt = int(b["gop_vao"])
                n = con.execute("UPDATE students SET %s=? WHERE %s=?" % (col, col), (tgt, cid)).rowcount
                con.execute("DELETE FROM %s WHERE id=?" % table, (cid,))
                log(con, None, "GỘP DANH MỤC", "%s #%d vào #%d (%d HV)" % (table, cid, tgt, n))
                con.commit()
                return dict(ok=True, gop=True, so=n)
            dup = con.execute("SELECT id FROM %s WHERE ten=?" % table, (ten,)).fetchone()
            if cid and dup and dup["id"] != cid:
                raise ValueError("GOP:%d:Tên '%s' đã có. Gộp 2 mục làm một (học viên chuyển theo)?" % (dup["id"], ten))
            args = [ten, clean(b.get("sdt")), clean(b.get("ghi_chu"))]
            if table == "teachers":
                xh = clean(b.get("xung_ho")).lower()
                args.append(xh if xh in XUNG_HO else "anh")
            xcol = ", xung_ho=?" if table == "teachers" else ""
            if cid:
                con.execute("UPDATE %s SET ten=?, sdt=?, ghi_chu=?%s, active=? WHERE id=?" % (table, xcol),
                            (*args, 1 if b.get("active", 1) else 0, cid))
            else:
                if dup:
                    raise ValueError("Tên '%s' đã có" % ten)
                if table == "runners" and not b.get("van_them"):
                    for r in con.execute("SELECT ten FROM runners"):
                        if person_key(r["ten"]) == person_key(ten):
                            raise ValueError("HOI:Đã có người chạy '%s' (có thể cùng một người). Vẫn thêm mới?" % r["ten"])
                if table == "teachers":
                    con.execute("INSERT INTO teachers(ten,sdt,ghi_chu,xung_ho,created_at) VALUES(?,?,?,?,?)", (*args, now_s()))
                else:
                    con.execute("INSERT INTO runners(ten,sdt,ghi_chu,created_at) VALUES(?,?,?,?)", (*args, now_s()))
            con.commit()
        finally:
            con.close()
    return dict(ok=True)

def delete_cat(table, cid):
    col = "gv_id" if table == "teachers" else "nth_id"
    with DB_LOCK:
        con = db()
        try:
            if table == "runners" and cid == mine_id(con):
                raise ValueError("Không xóa được MINE – đây là người chạy mặc định")
            n = con.execute("SELECT COUNT(*) c FROM students WHERE %s=? AND deleted=0" % col, (cid,)).fetchone()["c"]
            if n:
                raise ValueError("Còn %d học viên đang gán – hãy gộp/chuyển, hoặc bỏ tích 'đang làm' thay vì xóa" % n)
            con.execute("UPDATE students SET %s=NULL WHERE %s=?" % (col, col), (cid,))
            con.execute("DELETE FROM %s WHERE id=?" % table, (cid,))
            con.commit()
        finally:
            con.close()
    return dict(ok=True)

# ----------------------------------------------------------------- XUẤT EXCEL
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFILL = PatternFill("solid", fgColor="1F4E78")

def xsafe(v):
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v

def _hdr(ws, row, headers, widths=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HFILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        if widths:
            ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]

def export_xlsx(ids=None):
    S, st = compute_all()
    T = stats(S, st)
    if ids:
        ids = set(ids)
        S = [s for s in S if s["id"] in ids]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DANH SÁCH HỌC ONLINE"
    ws.cell(row=1, column=1, value="QUẢN LÝ DANH SÁCH HỌC ONLINE – " + st["ten_don_vi"]).font = Font(bold=True, size=14, color="1F4E78")
    ws.cell(row=2, column=1, value="Xuất ngày %s  |  Tổng: %d  |  Đã xong: %d  |  Đang học: %d  |  Chưa học: %d" % (
        today().strftime("%d/%m/%Y"), len(S), sum(s["trang_thai"] == ST_XONG for s in S),
        sum(s["trang_thai"] == ST_DANG for s in S), sum(s["trang_thai"] == ST_CHUA for s in S))).font = Font(italic=True)
    H = ["STT", "HỌ VÀ TÊN", "SỐ CCCD", "KHÓA HỌC", "GIÁO VIÊN", "SĐT GV", "ĐÃ ĐÓNG PHÍ", "TÌNH TRẠNG", "ĐÃ BÁO GV", "NGƯỜI THỰC HIỆN",
         "ĐÃ THANH TOÁN", "MÔN ĐẠT", "GIỜ ĐẠT", "ĐỊNH MỨC", "%", "GIỜ CÒN", "SỐ NGÀY CHẠY CÒN", "DỰ KIẾN XONG", "NGÀY LỌC", "GHI CHÚ"]
    _hdr(ws, 3, H, [5, 26, 15, 11, 22, 12, 10, 13, 10, 15, 12, 8, 9, 9, 7, 9, 9, 12, 11, 30])
    fills = {ST_XONG: PatternFill("solid", fgColor="E2EFDA"), ST_DANG: PatternFill("solid", fgColor="FFF2CC")}
    for i, s in enumerate(S, 1):
        has = bool(s["mons"])
        v = [i, s["ho_ten"], s["cccd"], s["khoa"], s.get("gv_ten") or "", s.get("gv_sdt") or "",
             "XONG" if s["phi_gv"] else "CHƯA", ST_LABEL[s["trang_thai"]], ("ĐÃ BÁO" if s["bao_gv"] else "CHƯA BÁO") if s["trang_thai"] == ST_XONG else "", s.get("nth_ten") or "", "ĐÃ TT" if s["nth_tt"] else "",
             "%d/%d" % (s["so_dat"], s["so_mon"]) if has else "", fmt_hm(s["tong_done"]) if has else "",
             fmt_hm(s["tong_norm"]) if has else "", s["pct"] / 100.0 if has else None, fmt_hm(s["con_lai"]) if has else "",
             s["so_ngay"] if s["so_ngay"] else "", ddmm(s["ngay_du_kien"]), ddmm(s["ngay_loc"]), s["ghi_chu"] or ""]
        for j, x in enumerate(v, 1):
            c = ws.cell(row=3 + i, column=j, value=xsafe(x))
            if j == 15 and x is not None: c.number_format = "0.0%"
            if j == 3: c.number_format = "@"
            if j == 8 and s["trang_thai"] in fills: c.fill = fills[s["trang_thai"]]
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = "A3:%s%d" % (get_column_letter(len(H)), 3 + len(S))
    w2 = wb.create_sheet("TIẾN ĐỘ TỪNG MÔN")
    allm = []
    for s in S:
        for m in s["mons"]:
            if m["short"] not in allm:
                allm.append(m["short"])
    _hdr(w2, 1, ["STT", "HỌ VÀ TÊN", "SỐ CCCD", "KHÓA"] + allm + ["NGÀY LỌC"], [5, 26, 15, 11] + [16] * len(allm) + [11])
    ok, no = PatternFill("solid", fgColor="E2EFDA"), PatternFill("solid", fgColor="FCE4D6")
    for i, s in enumerate(S, 1):
        for j, x in enumerate([i, s["ho_ten"], s["cccd"], s["khoa"]], 1):
            w2.cell(row=1 + i, column=j, value=x)
        for m in s["mons"]:
            c = w2.cell(row=1 + i, column=5 + allm.index(m["short"]), value="%s / %s %s" % (fmt_hm(m["done"]), fmt_hm(m["norm"]), "✔" if m["dat"] else ""))
            c.fill = ok if m["dat"] else no
        w2.cell(row=1 + i, column=5 + len(allm), value=ddmm(s["ngay_loc"]))
    w2.freeze_panes = "C2"
    w3 = wb.create_sheet("THỐNG KÊ")
    r = 1
    for title, key in (("THEO GIÁO VIÊN", "gv"), ("THEO NGƯỜI THỰC HIỆN", "nth"), ("THEO KHÓA", "khoa")):
        w3.cell(row=r, column=1, value=title).font = Font(bold=True, size=12, color="1F4E78")
        r += 1
        _hdr(w3, r, ["TÊN", "TỔNG HV", "CHƯA HỌC", "ĐANG HỌC", "ĐÃ XONG", "ĐÃ ĐÓNG PHÍ", "TIỀN ĐÃ THU", "CHƯA ĐÓNG", "TIỀN CHƯA ĐÓNG", "XONG CHƯA TRẢ NGƯỜI CHẠY"],
             [26, 9, 10, 10, 9, 11, 13, 10, 15, 14])
        for x in T[key]:
            r += 1
            for j, k in enumerate(["ten", "tong", "chua", "dang", "xong", "phi_ok", "tien_thu", "phi_no", "tien_no", "tt_no"], 1):
                c = w3.cell(row=r, column=j, value=xsafe(x[k]))
                if k.startswith("tien"): c.number_format = "#,##0"
        r += 2
    w3.cell(row=r, column=1, value="THEO THÁNG").font = Font(bold=True, size=12, color="1F4E78")
    r += 1
    _hdr(w3, r, ["THÁNG", "HV MỚI", "HV XONG", "SỐ HV ĐÓNG PHÍ", "TIỀN PHÍ THU", "SỐ HV ĐÃ TRẢ NGƯỜI CHẠY", "TIỀN TRẢ NGƯỜI CHẠY"])
    for x in T["thang"]:
        r += 1
        for j, k in enumerate(["thang", "moi", "xong", "phi_so", "phi_tien", "tt_so", "tt_tien"], 1):
            c = w3.cell(row=r, column=j, value=x[k])
            if "tien" in k: c.number_format = "#,##0"
    w4 = wb.create_sheet("DANH MỤC GV")
    _hdr(w4, 1, ["STT", "HỌ VÀ TÊN GIÁO VIÊN", "SỐ ĐIỆN THOẠI", "SỐ HV"], [5, 28, 15, 8])
    for i, t in enumerate(list_cat("teachers"), 1):
        for j, x in enumerate([i, t["ten"], t["sdt"], t["so_hv"]], 1):
            w4.cell(row=1 + i, column=j, value=x)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()

# ----------------------------------------------------------------- SAO LƯU
LAST_SIG = {"sig": None}

def data_signature():
    con = db()
    try:
        return tuple(con.execute(q).fetchone()[0] for q in (
            "SELECT MAX(id) FROM logs", "SELECT MAX(id) FROM imports", "SELECT MAX(updated_at) FROM students",
            "SELECT COUNT(*) FROM students", "SELECT group_concat(value) FROM settings",
            "SELECT group_concat(ten||sdt||active) FROM teachers", "SELECT group_concat(ten||sdt||active) FROM runners"))
    finally:
        con.close()

def check_db_file(path):
    t = sqlite3.connect(path)
    try:
        ok = t.execute("PRAGMA quick_check").fetchone()[0]
        n = t.execute("SELECT COUNT(*) FROM students WHERE deleted=0").fetchone()[0]
        return ok, n
    finally:
        t.close()

def do_backup(reason="auto", force=True):
    sig = data_signature()
    if not force and sig == LAST_SIG["sig"]:
        return dict(ok=True, skipped=True)
    st = get_settings()
    name = "hoc_online_%s_%s.db" % (dt.datetime.now().strftime("%Y%m%d_%H%M%S"), reason)
    dst = os.path.join(BACKUP_DIR, name)
    with DB_LOCK:
        con = sqlite3.connect(DB_PATH)
        try:
            con.execute("VACUUM INTO ?", (dst,))     # bản sao gọn, chế độ thường (không WAL) – chép lên Drive an toàn
        except sqlite3.OperationalError:
            out = sqlite3.connect(dst); con.backup(out); out.close()
        finally:
            con.close()
    ok, n = check_db_file(dst)
    if ok != "ok":
        os.remove(dst)
        raise RuntimeError("Bản sao lưu không qua kiểm tra toàn vẹn: %s" % ok)
    LAST_SIG["sig"] = sig
    keep_n = int(st["backup_giu"])
    extra = (st["backup_extra_dir"] or "").strip()
    extra_ok = None
    if extra:
        try:
            os.makedirs(extra, exist_ok=True)
            shutil.copy2(dst, os.path.join(extra, name))
            shutil.copy2(dst, os.path.join(extra, "hoc_online_MOI_NHAT.db"))
            prune(extra, keep_n)
            extra_ok = True
        except Exception as e:
            extra_ok = str(e)
    prune(BACKUP_DIR, keep_n)
    return dict(ok=True, file=name, hv=n, extra=extra_ok)

TRASH = "_thung_rac"
TRASH_DAYS = 30

def prune(folder, keep_n=10):
    """Giữ N bản mới nhất; bản cũ hơn chuyển vào _thung_rac, quá 30 ngày trong thùng rác thì xóa hẳn."""
    files = sorted(f for f in os.listdir(folder) if re.match(r"^hoc_online_\d{8}_\d{6}_.*\.db$", f))
    trash = os.path.join(folder, TRASH)
    old = files[:-keep_n] if len(files) > keep_n else []
    if old:
        os.makedirs(trash, exist_ok=True)
    for f in old:
        try:
            dst = os.path.join(trash, f)
            shutil.move(os.path.join(folder, f), dst)
            os.utime(dst, None)                       # mốc tính 30 ngày = lúc vào thùng rác
        except OSError:
            pass
    if os.path.isdir(trash):
        lim = time.time() - TRASH_DAYS * 86400
        for f in os.listdir(trash):
            p = os.path.join(trash, f)
            try:
                if os.path.getmtime(p) < lim:
                    os.remove(p)
            except OSError:
                pass

def list_backups():
    out = []
    for folder, trash in ((BACKUP_DIR, 0), (os.path.join(BACKUP_DIR, TRASH), 1)):
        if not os.path.isdir(folder):
            continue
        for f in sorted(os.listdir(folder), reverse=True):
            if f.endswith(".db"):
                p = os.path.join(folder, f)
                mt = os.path.getmtime(p)
                d = dict(file=f, size=os.path.getsize(p), trash=trash,
                         time=dt.datetime.fromtimestamp(mt).strftime("%d/%m/%Y %H:%M:%S") if not trash else f[11:19][6:8] + "/" + f[11:19][4:6] + "/" + f[11:15] + " " + f[20:22] + ":" + f[22:24])
                if trash:
                    d["con_ngay"] = max(0, TRASH_DAYS - int((time.time() - mt) // 86400))
                out.append(d)
    return out

def backup_path(fname):
    f = os.path.basename(fname)
    for folder in (BACKUP_DIR, os.path.join(BACKUP_DIR, TRASH)):
        p = os.path.join(folder, f)
        if os.path.isfile(p):
            return p
    raise ValueError("Không thấy file sao lưu")

def backup_info(fname):
    p = backup_path(fname)
    try:
        ok, n = check_db_file(p)
    except Exception as e:
        raise ValueError("File sao lưu hỏng: %s" % e)
    con = db()
    cur = con.execute("SELECT COUNT(*) FROM students WHERE deleted=0").fetchone()[0]
    con.close()
    return dict(ok=ok, hv_backup=n, hv_hien_tai=cur)

def restore_backup(fname):
    info = backup_info(fname)
    if info["ok"] != "ok":
        raise ValueError("File sao lưu không toàn vẹn – không khôi phục")
    p = backup_path(fname)
    tmp = p + ".tmp"
    shutil.copy2(p, tmp)                      # chép ra trước: bản sao lưu mới có thể đẩy file này vào thùng rác
    do_backup("truoc_khoi_phuc")
    p = tmp
    with DB_LOCK:
        src = sqlite3.connect(p)
        out = sqlite3.connect(DB_PATH)
        src.backup(out)
        out.close()
        src.close()
    try:
        os.remove(tmp)
    except OSError:
        pass
    LAST_SIG["sig"] = None
    init_db()
    return dict(ok=True, hv=info["hv_backup"])

def backup_loop():
    last = time.time()
    while True:
        time.sleep(15)
        try:
            if time.time() - last >= int(get_settings()["backup_interval_m"]) * 60:
                last = time.time()
                do_backup("auto", force=False)        # chỉ sao lưu khi dữ liệu có thay đổi
            with PREVIEW_LOCK:
                for k in [k for k, v in PREVIEW.items() if time.time() - v["created"] > 7200]:
                    PREVIEW.pop(k, None)
        except Exception:
            traceback.print_exc()

# ----------------------------------------------------------------- CẬP NHẬT TỪ GITHUB
UPDATE_STATE = {"checked": None, "latest": None, "notes": "", "error": ""}
SRV = {"srv": None}

def ver_tuple(v):
    return tuple(int(x) for x in re.findall(r"\d+", str(v))[:4])

def http_get(url, timeout=20):
    req = urllib.request.Request(url, headers={"User-Agent": "hoc-online-oto/" + VERSION, "Cache-Control": "no-cache"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()

def check_update():
    url = "https://raw.githubusercontent.com/%s/%s/version.json?t=%d" % (GITHUB_REPO, GITHUB_BRANCH, int(time.time()))
    try:
        info = json.loads(http_get(url, 10).decode("utf-8"))
        UPDATE_STATE.update(checked=now_s(), latest=info.get("version"), notes=info.get("notes", ""),
                            date=info.get("date", ""), error="")
    except Exception as e:
        msg = "Chưa có thông tin phiên bản trên GitHub" if getattr(e, "code", None) == 404 else "Không kết nối được GitHub – kiểm tra mạng internet (%s)" % e.__class__.__name__
        UPDATE_STATE.update(checked=now_s(), error=msg)
    UPDATE_STATE["current"] = VERSION
    UPDATE_STATE["co_ban_moi"] = bool(UPDATE_STATE.get("latest") and ver_tuple(UPDATE_STATE["latest"]) > ver_tuple(VERSION))
    UPDATE_STATE["repo"] = "https://github.com/" + GITHUB_REPO
    return dict(UPDATE_STATE)

SKIP_TOP = ("data", "backups", ".git", ".github")

def apply_update():
    """Tải bản mới nhất từ GitHub, sao lưu dữ liệu + code cũ, chép đè code (không đụng data/backups)."""
    import zipfile
    data = http_get("https://codeload.github.com/%s/zip/refs/heads/%s" % (GITHUB_REPO, GITHUB_BRANCH), 120)
    zf = zipfile.ZipFile(io.BytesIO(data))
    names = zf.namelist()
    top = names[0].split("/")[0] + "/"
    if top + "app/server.py" not in names:
        raise ValueError("Gói tải về không đúng cấu trúc tool – không cập nhật")
    m = re.search(rb'VERSION = "([^"]+)"', zf.read(top + "app/server.py"))
    new_ver = m.group(1).decode() if m else "?"
    do_backup("truoc_cap_nhat")
    code_bak = os.path.join(BACKUP_DIR, "code_v%s_%s" % (VERSION, dt.datetime.now().strftime("%Y%m%d_%H%M%S")))
    shutil.copytree(BASE, os.path.join(code_bak, "app"), ignore=shutil.ignore_patterns("__pycache__"))
    n = 0
    root_real = os.path.realpath(ROOT)
    for nm in names:
        if not nm.startswith(top) or nm.endswith("/"):
            continue
        rel = nm[len(top):]
        if not rel or rel.split("/")[0] in SKIP_TOP or ".." in rel.split("/"):
            continue
        dst = os.path.realpath(os.path.join(ROOT, *rel.split("/")))
        if not dst.startswith(root_real + os.sep):
            continue
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        with open(dst, "wb") as f:
            f.write(zf.read(nm))
        n += 1
    return dict(ok=True, version=new_ver, so_file=n, code_bak=code_bak)

def restart_later():
    def go():
        time.sleep(1.0)
        try:
            if SRV["srv"]:
                SRV["srv"].shutdown()
                SRV["srv"].server_close()
        except Exception:
            pass
        kw = {"cwd": ROOT}
        if os.name == "nt":
            kw["creationflags"] = getattr(subprocess, "CREATE_NEW_CONSOLE", 0x10)
        subprocess.Popen([sys.executable, os.path.join(BASE, "server.py"), "--no-browser"], **kw)
        os._exit(0)
    threading.Thread(target=go, daemon=True).start()

def save_settings(b):
    vals = {k: validate_setting(k, v) for k, v in b.items() if k in SETTINGS_SPEC}   # sai 1 ô -> báo lỗi, không lưu gì
    with DB_LOCK:
        con = db()
        for k, v in vals.items():
            con.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (k, v))
        con.commit()
        con.close()
    return dict(ok=True, settings=get_settings())

# ----------------------------------------------------------------- HTTP
class H(BaseHTTPRequestHandler):
    server_version = "HocOnline/" + VERSION

    def log_message(self, fmt, *args):
        pass

    def _send(self, code, body, ctype="application/json; charset=utf-8", extra=None):
        if isinstance(body, (dict, list)):
            body = json.dumps(body, ensure_ascii=False, default=str).encode("utf-8")
        elif isinstance(body, str):
            body = body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def _body(self):
        n = int(self.headers.get("Content-Length") or 0)
        if n > 50 * 1024 * 1024:
            raise ValueError("File quá lớn")
        return self.rfile.read(n) if n else b""

    def do_GET(self):
        u = urlparse(self.path)
        q = {k: v[0] for k, v in parse_qs(u.query).items()}
        try:
            p = u.path
            if p in ("/", "/index.html"):
                with open(os.path.join(STATIC_DIR, "index.html"), "rb") as f:
                    return self._send(200, f.read(), "text/html; charset=utf-8")
            if p == "/api/all":
                S, st = compute_all()
                return self._send(200, dict(students=S, settings=st, rem=reminders(S, st), teachers=list_cat("teachers"),
                                            runners=list_cat("runners"), today=today().isoformat(), version=VERSION, version_date=VERSION_DATE,
                                            update=dict(co_ban_moi=UPDATE_STATE.get("co_ban_moi", False), latest=UPDATE_STATE.get("latest"))))
            if p == "/api/students":
                return self._send(200, compute_all(q.get("deleted") == "1")[0])
            if p == "/api/student":
                return self._send(200, student_detail(int(q["id"])))
            if p == "/api/stats":
                S, st = compute_all()
                return self._send(200, dict(stats(S, st), rem=reminders(S, st)))
            if p == "/api/settings":
                return self._send(200, get_settings())
            if p == "/api/imports":
                con = db()
                r = [dict(x) for x in con.execute("SELECT id,ten_file,loai,khoa,ngay_loc,so_dong,so_khop,so_them,so_xong,ghi_chu,created_at,undone FROM imports ORDER BY id DESC LIMIT 300")]
                con.close()
                return self._send(200, r)
            if p == "/api/backups":
                return self._send(200, list_backups())
            if p == "/api/update":
                return self._send(200, dict(UPDATE_STATE, current=VERSION, repo="https://github.com/" + GITHUB_REPO))
            if p == "/api/backup_info":
                return self._send(200, backup_info(q.get("file", "")))
            if p == "/api/info":
                return self._send(200, dict(app=APP_ID, name=APP_NAME, version=VERSION, db=DB_PATH, backup_dir=BACKUP_DIR, today=today().isoformat()))
            if p == "/api/export.xlsx":
                ids = [int(x) for x in q.get("ids", "").split(",") if x.strip().isdigit()]
                data = export_xlsx(ids or None)
                fn = "HOC_ONLINE_%s.xlsx" % dt.datetime.now().strftime("%Y%m%d_%H%M")
                return self._send(200, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                  {"Content-Disposition": 'attachment; filename="%s"' % fn})
            return self._send(404, {"error": "Không có"})
        except (ValueError, KeyError) as e:
            return self._send(400, {"error": str(e)})
        except Exception as e:
            traceback.print_exc()
            return self._send(500, {"error": "Lỗi: %s" % e})

    def do_POST(self):
        u = urlparse(self.path)
        q = {k: v[0] for k, v in parse_qs(u.query).items()}
        p = u.path
        try:
            if p == "/api/import/preview":
                return self._send(200, preview_file(unquote(q.get("name", "file.xlsx")), self._body()))
            raw = self._body()
            try:
                b = json.loads(raw.decode("utf-8")) if raw else {}
            except ValueError:
                raise ValueError("Dữ liệu gửi lên không hợp lệ")
            if p == "/api/import/commit":
                return self._send(200, commit_report(b) if b.get("kind") == "report" else commit_ql(b))
            if p == "/api/import/undo":
                return self._send(200, undo_import(int(b["id"])))
            if p == "/api/students/update":
                return self._send(200, update_students([int(i) for i in b["ids"]], b["fields"], b.get("note", "")))
            if p == "/api/students/add":
                return self._send(200, add_students(b.get("list") or [b], b.get("common") or {}))
            if p == "/api/students/delete":
                return self._send(200, delete_students([int(i) for i in b["ids"]], bool(b.get("restore"))))
            if p == "/api/students/purge":
                return self._send(200, purge_students([int(i) for i in b["ids"]]))
            if p in ("/api/teachers/save", "/api/runners/save"):
                return self._send(200, save_cat("teachers" if "teachers" in p else "runners", b))
            if p in ("/api/teachers/delete", "/api/runners/delete"):
                return self._send(200, delete_cat("teachers" if "teachers" in p else "runners", int(b["id"])))
            if p == "/api/settings":
                return self._send(200, save_settings(b))
            if p == "/api/backup":
                return self._send(200, do_backup("tay"))
            if p == "/api/restore":
                return self._send(200, restore_backup(b["file"]))
            if p == "/api/backup/untrash":
                src = backup_path(b["file"])
                shutil.move(src, os.path.join(BACKUP_DIR, os.path.basename(src)))
                return self._send(200, {"ok": True})
            if p == "/api/update/check":
                return self._send(200, check_update())
            if p == "/api/update/apply":
                r = apply_update()
                restart_later()
                return self._send(200, r)
            return self._send(404, {"error": "Không có"})
        except (ValueError, KeyError) as e:
            return self._send(400, {"error": str(e) if isinstance(e, ValueError) else "Thiếu dữ liệu: %s" % e})
        except Exception as e:
            traceback.print_exc()
            return self._send(500, {"error": "Lỗi: %s" % e})

def is_our_tool(port):
    try:
        with urllib.request.urlopen("http://127.0.0.1:%d/api/info" % port, timeout=1.5) as r:
            return json.loads(r.read().decode("utf-8")).get("app") == APP_ID
    except Exception:
        return False

def port_free(port):
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        try:
            s.bind(("127.0.0.1", port))
            return True
        except OSError:
            return False

def open_browser(url):
    if "--no-browser" in sys.argv:
        return
    for env in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA"):
        base = os.environ.get(env)
        if base:
            c = os.path.join(base, "Google", "Chrome", "Application", "chrome.exe")
            if os.path.isfile(c):
                try:
                    subprocess.Popen([c, "--app=" + url])
                    return
                except Exception:
                    pass
    webbrowser.open(url)

def cli_update():
    print("Dang tai ban moi nhat tu GitHub: %s ..." % GITHUB_REPO)
    init_db()
    r = apply_update()
    print("Da cap nhat len v%s (%d file). Code cu luu o: %s" % (r["version"], r["so_file"], r["code_bak"]))

def main():
    if "--update" in sys.argv:
        try:
            cli_update()
        except Exception as e:
            print("LOI cap nhat:", e)
        return
    port = PORT
    for _ in range(10):
        if is_our_tool(port):
            print("Tool dang chay san -> mo trinh duyet")
            open_browser("http://127.0.0.1:%d" % port)
            return
        if port_free(port):
            break
        port += 1
    url = "http://127.0.0.1:%d" % port
    init_db()
    try:
        do_backup("khoi_dong")
    except Exception:
        traceback.print_exc()
    threading.Thread(target=backup_loop, daemon=True).start()
    srv = ThreadingHTTPServer(("127.0.0.1", port), H)
    SRV["srv"] = srv
    if get_settings()["tu_kiem_tra_cap_nhat"] == "1":
        threading.Thread(target=check_update, daemon=True).start()
    print("=" * 62)
    print(" %s  v%s" % (strip_accents(APP_NAME), VERSION))
    print(" Dang chay tai : %s" % url)
    print(" Du lieu       : %s" % DB_PATH)
    print(" Sao luu       : %s" % BACKUP_DIR)
    print(" KHONG TAT cua so nay khi dang dung tool (thu nho lai duoc).")
    print("=" * 62)
    threading.Timer(1.0, lambda: open_browser(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        try:
            do_backup("tat_tool", force=False)
        except Exception:
            pass

if __name__ == "__main__":
    main()
